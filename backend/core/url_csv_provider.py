import csv
import hashlib
import io
import os
import re
import time
import zipfile
import gzip
import ftplib
import tempfile
import threading
import json
import sqlite3
from collections import Counter
from contextlib import closing
from urllib.parse import urlparse
from urllib.parse import unquote
import requests

from config_path import get_config_dir
from price_names import humanize_price_name

try:
    csv.field_size_limit(1024 * 1024 * 1024)
except OverflowError:
    csv.field_size_limit(256 * 1024 * 1024)


class UrlCsvProvider:
    """Поставщик цен по CSV-файлам, доступным по ссылке."""

    def __init__(
        self,
        urls=None,
        name="Freno CSV",
        timeout=20,
        column_map=None,
        default_days=0,
        url_profiles=None,
        cache_hours=24,
        runtime_refresh=True,
    ):
        self.urls = [str(item).strip() for item in (urls or []) if str(item).strip()]
        self.name = name or "Freno CSV"
        self.timeout = timeout
        self.column_map = column_map or {}
        self.default_days = default_days or 0
        self.url_profiles = url_profiles or {}
        self.cache_hours = max(1, int(cache_hours or 24))
        self.runtime_refresh = bool(runtime_refresh)
        self.cache_dir = os.path.join(get_config_dir(), "csv_cache")
        self.last_message = ""
        self.debug_log = []
        self._sqlite_ready = False
        self._index_lock = threading.Lock()

    def get_prices(self, article):
        self.last_message = ""
        if not self.urls:
            self.last_message = "Нет ссылок на CSV"
            return []
        target_article = self._normalize_article(article)
        if not target_article:
            self.last_message = "Пустой артикул"
            return []

        if not self._ensure_sqlite_index(
            local_only=not self.runtime_refresh,
            build_missing=self.runtime_refresh,
        ):
            return []
        active_urls = [
            url for url in self.urls
            if (self._profile_for_url(url) or {}).get("enabled", True) is not False
        ]
        if not active_urls:
            self.last_message = "Нет включённых CSV"
            return []
        placeholders = ",".join("?" for _ in active_urls)
        query = f"""
            SELECT provider, brand, article, price, days, quantity, logo,
                   warehouse, name, multiplicity, is_cross, source_url,
                   source_brand, source_code, delivery_total_hours
            FROM offers
            WHERE article_key = ? AND source_url IN ({placeholders})
        """
        try:
            with closing(sqlite3.connect(self._index_path(), timeout=30)) as connection:
                rows = connection.execute(query, [target_article, *active_urls]).fetchall()
        except sqlite3.Error as exc:
            self.last_message = f"{self.name}: ошибка локальной базы: {exc}"
            return []
        matches = [
            {
                "provider": self._display_source_name(row[0]), "brand": row[1], "article": row[2],
                "price": row[3], "days": row[4], "quantity": row[5],
                "logo": row[6], "warehouse": row[7], "name": row[8],
                "multiplicity": row[9], "is_cross": bool(row[10]),
                "source_url": row[11], "source_brand": row[12],
                "source_code": row[13], "delivery_total_hours": row[14],
            }
            for row in rows
        ]
        if matches:
            self.last_message = ""
            return matches

        self.last_message = f"{self.name}: нет данных для {article}"
        return []

    def get_prices_many(self, articles):
        self.last_message = ""
        if not self.urls:
            self.last_message = "Нет ссылок на CSV"
            return {}
        article_keys = []
        requested_aliases = {}
        seen_articles = set()
        for article in articles or []:
            key = self._normalize_article(article)
            alias = str(article or "").strip()
            canonical_alias = re.sub(r"[^A-Z0-9]+", "", alias.upper())
            if key and key not in seen_articles:
                seen_articles.add(key)
                article_keys.append(key)
            if key and alias:
                requested_aliases.setdefault(key, set()).add(alias)
            if key and canonical_alias:
                requested_aliases.setdefault(key, set()).add(canonical_alias)
        if not article_keys:
            self.last_message = "Пустой список артикулов"
            return {}

        if not self._ensure_sqlite_index(
            local_only=not self.runtime_refresh,
            build_missing=self.runtime_refresh,
        ):
            return {}
        active_urls = [
            url for url in self.urls
            if (self._profile_for_url(url) or {}).get("enabled", True) is not False
        ]
        if not active_urls:
            self.last_message = "Нет включённых CSV"
            return {}

        rows_by_key = {key: [] for key in article_keys}
        url_placeholders = ",".join("?" for _ in active_urls)
        try:
            with closing(sqlite3.connect(self._index_path(), timeout=30)) as connection:
                for chunk in self._chunks(article_keys, 400):
                    article_placeholders = ",".join("?" for _ in chunk)
                    query = f"""
                        SELECT article_key, {self._select_columns()}
                        FROM offers
                        WHERE article_key IN ({article_placeholders})
                          AND source_url IN ({url_placeholders})
                        ORDER BY article_key ASC, days ASC, price ASC
                    """
                    rows = connection.execute(query, [*chunk, *active_urls]).fetchall()
                    for row in rows:
                        rows_by_key.setdefault(row[0], []).append(self._item_from_index_row(row[1:]))
        except sqlite3.Error as exc:
            self.last_message = f"{self.name}: ошибка локальной базы: {exc}"
            return {}

        result = {}
        for key, rows in rows_by_key.items():
            result[key] = rows
            for alias in requested_aliases.get(key, set()):
                result[alias] = rows
        found_count = sum(len(items) for key, items in rows_by_key.items())
        if found_count:
            self.last_message = ""
            return result
        self.last_message = f"{self.name}: нет данных для списка кроссов"
        return result

    def search_local_cache(self, query_text, limit=300):
        self.last_message = ""
        query_text = str(query_text or "").strip()
        if not query_text:
            self.last_message = "Пустой запрос"
            return []
        if not self.urls:
            self.last_message = "Нет ссылок на прайсы"
            return []

        self._ensure_sqlite_index(local_only=True, build_missing=False)
        active_urls = [
            url for url in self.urls
            if (self._profile_for_url(url) or {}).get("enabled", True) is not False
        ]
        if not active_urls:
            self.last_message = "Нет включённых прайсов"
            return []

        limit = max(1, min(1000, int(limit or 300)))
        rows = []
        seen = set()
        article_key = self._normalize_article(query_text)
        groups = self._search_groups(query_text)

        try:
            with closing(sqlite3.connect(self._index_path(), timeout=30)) as connection:
                active_urls = self._indexed_active_urls(connection, active_urls)
                if not active_urls:
                    self.last_message = f"{self.name}: локальная база не готова, обновите прайсы"
                    return []
                placeholders = ",".join("?" for _ in active_urls)
                if len(article_key) >= 3:
                    exact_sql = f"""
                        SELECT {self._select_columns()}
                        FROM offers
                        WHERE article_key = ? AND source_url IN ({placeholders})
                        ORDER BY days ASC, price ASC
                        LIMIT ?
                    """
                    for row in connection.execute(exact_sql, [article_key, *active_urls, limit]).fetchall():
                        marker = self._row_marker(row)
                        if marker not in seen:
                            seen.add(marker)
                            rows.append(row)
                remaining = limit - len(rows)
                if groups and remaining > 0:
                    clauses = []
                    params = []
                    for terms in groups[:6]:
                        term_clauses = []
                        for term in terms[:5]:
                            term_clauses.append("search_text LIKE ?")
                            params.append(f"%{term}%")
                        if term_clauses:
                            clauses.append("(" + " AND ".join(term_clauses) + ")")
                    if clauses:
                        text_sql = f"""
                            SELECT {self._select_columns()}
                            FROM offers
                            WHERE source_url IN ({placeholders})
                              AND ({' OR '.join(clauses)})
                            ORDER BY days ASC, price ASC
                            LIMIT ?
                        """
                        for row in connection.execute(text_sql, [*active_urls, *params, remaining]).fetchall():
                            marker = self._row_marker(row)
                            if marker not in seen:
                                seen.add(marker)
                                rows.append(row)
        except sqlite3.Error as exc:
            self.last_message = f"{self.name}: ошибка локальной базы: {exc}"
            return []

        result = [self._item_from_index_row(row) for row in rows]
        if not result:
            self.last_message = f"{self.name}: в локальной базе ничего не найдено"
        return result

    def suggest_articles(self, query_text, limit=10):
        self.last_message = ""
        query_key = self._normalize_article(query_text)
        if len(query_key) < 3:
            return []
        if not self.urls:
            self.last_message = "Нет ссылок на прайсы"
            return []

        if not self._ensure_sqlite_index(local_only=True, build_missing=False):
            return []
        active_urls = [
            url for url in self.urls
            if (self._profile_for_url(url) or {}).get("enabled", True) is not False
        ]
        if not active_urls:
            self.last_message = "Нет включённых прайсов"
            return []

        limit = max(1, min(50, int(limit or 10)))
        try:
            with closing(sqlite3.connect(self._index_path(), timeout=0.2)) as connection:
                active_urls = self._indexed_active_urls(connection, active_urls)
                if not active_urls:
                    self.last_message = f"{self.name}: локальная база не готова, обновите прайсы"
                    return []
                placeholders = ",".join("?" for _ in active_urls)
                rows = connection.execute(
                    f"""
                    SELECT article_key,
                           MIN(article) AS display_article,
                           GROUP_CONCAT(DISTINCT brand) AS brands,
                           MIN(name) AS display_name,
                           COUNT(*) AS offer_count,
                           MIN(CASE WHEN price > 0 THEN price END) AS min_price
                    FROM offers
                    WHERE article_key LIKE ?
                      AND source_url IN ({placeholders})
                    GROUP BY article_key
                    ORDER BY article_key ASC, offer_count DESC, min_price ASC
                    LIMIT ?
                    """,
                    [f"{query_key}%", *active_urls, limit],
                ).fetchall()
        except sqlite3.Error as exc:
            self.last_message = f"{self.name}: ошибка локальной базы: {exc}"
            return []

        suggestions = []
        for _article_key, article, brands, name, offer_count, min_price in rows:
            suggestions.append({
                "article": article or _article_key,
                "brand": self._short_brand_list(brands),
                "name": name or "",
                "source": "локальный прайс",
                "offer_count": int(offer_count or 0),
                "min_price": min_price,
            })
        if not suggestions:
            self.last_message = f"{self.name}: подсказки не найдены"
        return suggestions

    @staticmethod
    def _indexed_active_urls(connection, active_urls):
        active_urls = [url for url in active_urls if url]
        if not active_urls:
            return []
        placeholders = ",".join("?" for _ in active_urls)
        rows = connection.execute(
            f"SELECT source_url FROM indexed_sources WHERE source_url IN ({placeholders})",
            active_urls,
        ).fetchall()
        indexed = {row[0] for row in rows}
        return [url for url in active_urls if url in indexed]

    @staticmethod
    def _short_brand_list(value, limit=3):
        brands = []
        seen = set()
        for raw in str(value or "").split(","):
            brand = " ".join(str(raw or "").split())
            if not brand:
                continue
            key = re.sub(r"[^A-Z0-9]+", "", brand.upper())
            if key in seen:
                continue
            seen.add(key)
            brands.append(brand)
        if len(brands) <= limit:
            return "/".join(brands)
        return "/".join(brands[:limit]) + f" +{len(brands) - limit}"

    def get_brand_candidates(self, article):
        self.last_message = ""
        if not self.urls:
            self.last_message = "Нет ссылок на CSV"
            return []
        target_article = self._normalize_article(article)
        if not target_article:
            self.last_message = "Пустой артикул"
            return []

        if not self._ensure_sqlite_index(
            local_only=not self.runtime_refresh,
            build_missing=self.runtime_refresh,
        ):
            return []
        active_urls = [
            url for url in self.urls
            if (self._profile_for_url(url) or {}).get("enabled", True) is not False
        ]
        if not active_urls:
            self.last_message = "Нет включённых CSV"
            return []
        placeholders = ",".join("?" for _ in active_urls)
        query = f"""
            SELECT brand, article, name, provider, source_url
            FROM offers
            WHERE article_key = ? AND source_url IN ({placeholders})
        """
        try:
            with closing(sqlite3.connect(self._index_path(), timeout=30)) as connection:
                rows = connection.execute(query, [target_article, *active_urls]).fetchall()
        except sqlite3.Error as exc:
            self.last_message = f"{self.name}: ошибка локальной базы: {exc}"
            return []

        result = []
        seen = set()
        for brand, item_article, name, provider_name, source_url in rows:
            brand = str(brand or "").strip()
            if not brand:
                continue
            marker = (
                brand.upper(),
                self._normalize_article(item_article),
                str(name or "").upper(),
                str(provider_name or ""),
            )
            if marker in seen:
                continue
            seen.add(marker)
            result.append({
                "brand": brand,
                "article": item_article or article,
                "name": name or "",
                "provider": provider_name or self.name,
                "source": "CSV локальный прайс",
                "source_url": source_url,
            })
        if not result:
            self.last_message = f"{self.name}: бренды не найдены для {article}"
        return result

    def get_brands(self, article):
        brands = []
        seen = set()
        for item in self.get_brand_candidates(article):
            brand = str(item.get("brand") or "").strip()
            key = brand.upper()
            if brand and key not in seen:
                seen.add(key)
                brands.append(brand)
        return brands

    def _index_path(self):
        return os.path.join(self.cache_dir, "article_index.sqlite3")

    @staticmethod
    def _select_columns():
        return (
            "provider, brand, article, price, days, quantity, logo, "
            "warehouse, name, multiplicity, is_cross, source_url, "
            "source_brand, source_code, delivery_total_hours"
        )

    def _item_from_index_row(self, row):
        return {
            "provider": self._display_source_name(row[0]), "brand": row[1], "article": row[2],
            "price": row[3], "days": row[4], "quantity": row[5],
            "logo": row[6], "warehouse": row[7], "name": row[8],
            "multiplicity": row[9], "is_cross": bool(row[10]),
            "source_url": row[11], "source_brand": row[12],
            "source_code": row[13], "delivery_total_hours": row[14],
        }

    @staticmethod
    def _row_marker(row):
        return (
            str(row[0] or ""),
            str(row[1] or ""),
            str(row[2] or ""),
            float(row[3] or 0),
            str(row[7] or ""),
            str(row[11] or ""),
        )

    @staticmethod
    def _chunks(items, size):
        size = max(1, int(size or 1))
        for index in range(0, len(items), size):
            yield items[index:index + size]

    def _ensure_sqlite_schema(self, connection):
        connection.execute("""
            CREATE TABLE IF NOT EXISTS offers (
                article_key TEXT NOT NULL,
                source_url TEXT NOT NULL,
                provider TEXT, brand TEXT, article TEXT, price REAL,
                days INTEGER, quantity TEXT, logo TEXT, warehouse TEXT,
                name TEXT, multiplicity INTEGER, is_cross INTEGER,
                source_brand TEXT, source_code TEXT,
                delivery_total_hours INTEGER,
                search_text TEXT
            )
        """)
        columns = {
            row[1] for row in connection.execute("PRAGMA table_info(offers)").fetchall()
        }
        if "delivery_total_hours" not in columns:
            connection.execute("ALTER TABLE offers ADD COLUMN delivery_total_hours INTEGER")
        if "search_text" not in columns:
            connection.execute("ALTER TABLE offers ADD COLUMN search_text TEXT")
        connection.execute("""
            CREATE INDEX IF NOT EXISTS idx_offers_article
            ON offers(article_key)
        """)
        connection.execute("""
            CREATE INDEX IF NOT EXISTS idx_offers_source
            ON offers(source_url)
        """)
        connection.execute("""
            CREATE INDEX IF NOT EXISTS idx_offers_article_source
            ON offers(article_key, source_url)
        """)
        connection.execute("""
            CREATE INDEX IF NOT EXISTS idx_offers_search_text
            ON offers(search_text)
        """)
        connection.execute("""
            CREATE TABLE IF NOT EXISTS indexed_sources (
                source_url TEXT PRIMARY KEY,
                signature TEXT NOT NULL,
                row_count INTEGER NOT NULL DEFAULT 0
            )
        """)

    def _ensure_sqlite_index(self, local_only=False, build_missing=True):
        if self._sqlite_ready:
            return True
        acquired = self._index_lock.acquire(blocking=bool(build_missing))
        if not acquired:
            self.last_message = f"{self.name}: локальная база занята обновлением прайсов"
            return False
        try:
            if self._sqlite_ready:
                return True
            os.makedirs(self.cache_dir, exist_ok=True)
            with closing(sqlite3.connect(self._index_path(), timeout=60)) as connection:
                self._ensure_sqlite_schema(connection)
                indexed = dict(connection.execute(
                    "SELECT source_url, signature FROM indexed_sources"
                ).fetchall())
                connection.commit()
            for url in self.urls:
                profile = self._profile_for_url(url) or {}
                if profile.get("enabled", True) is False:
                    continue
                if indexed.get(url) == self._source_signature(url):
                    continue
                if not build_missing:
                    continue
                rows = self._load_cached_rows(url) if local_only else self._load_rows(url)
                if rows is None:
                    continue
                self._replace_source_index(url, rows)
            self._sqlite_ready = True
            return True
        finally:
            self._index_lock.release()

    def _load_cached_rows(self, url):
        cache_path = self._cache_path(url)
        if not os.path.exists(cache_path):
            self.last_message = f"{self.name}: локальный кэш не найден, обновите прайсы"
            return None
        try:
            with open(cache_path, "rb") as file:
                content = file.read()
            if urlparse(str(url)).path.lower().endswith(".xlsx"):
                return self._parse_xlsx(content)
            return self._parse_csv(self._decode_bytes(content), self._profile_for_url(url))
        except Exception as exc:
            self.last_message = f"{self.name}: не удалось прочитать локальный кэш: {exc}"
            return None

    def _source_signature(self, url):
        path = self._cache_path(url)
        try:
            stat = os.stat(path)
            version = (stat.st_mtime_ns, stat.st_size)
        except OSError:
            version = (0, 0)
        profile = self._profile_for_url(url) or {}
        index_profile = {
            key: profile.get(key)
            for key in ("enabled", "name", "has_header", "column_map", "default_days")
            if key in profile
        }
        raw = ("v3-csv-dialect-delivery", url, version, json.dumps(index_profile, sort_keys=True, ensure_ascii=False))
        return hashlib.sha256(repr(raw).encode("utf-8")).hexdigest()

    def _replace_source_index(self, url, rows):
        insert_sql = """
            INSERT INTO offers (
                article_key, source_url, provider, brand, article, price,
                days, quantity, logo, warehouse, name, multiplicity,
                is_cross, source_brand, source_code, delivery_total_hours,
                search_text
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """
        os.makedirs(self.cache_dir, exist_ok=True)
        count = 0
        with closing(sqlite3.connect(self._index_path(), timeout=120)) as connection:
            connection.execute("PRAGMA journal_mode=WAL")
            connection.execute("PRAGMA synchronous=NORMAL")
            self._ensure_sqlite_schema(connection)
            connection.execute("DELETE FROM offers WHERE source_url = ?", (url,))
            batch = []
            for row in rows:
                item = self._normalize_row(row, url)
                if not item:
                    continue
                article_key = self._normalize_article(item.get("article"))
                if not article_key:
                    continue
                search_text = self._index_search_text(item)
                batch.append((
                    article_key, url, item.get("provider", ""), item.get("brand", ""),
                    item.get("article", ""), float(item.get("price", 0) or 0),
                    int(item.get("days", 0) or 0), str(item.get("quantity", "0")),
                    item.get("logo", ""), item.get("warehouse", ""), item.get("name", ""),
                    int(item.get("multiplicity", 1) or 1), int(bool(item.get("is_cross"))),
                    item.get("source_brand", ""), item.get("source_code", ""),
                    int(item.get("delivery_total_hours", int(item.get("days", 0) or 0) * 24) or 0),
                    search_text,
                ))
                if len(batch) >= 5000:
                    connection.executemany(insert_sql, batch)
                    count += len(batch)
                    batch.clear()
            if batch:
                connection.executemany(insert_sql, batch)
                count += len(batch)
            connection.execute("""
                INSERT INTO indexed_sources(source_url, signature, row_count)
                VALUES (?, ?, ?)
                ON CONFLICT(source_url) DO UPDATE SET
                    signature=excluded.signature, row_count=excluded.row_count
            """, (url, self._source_signature(url), count))
            connection.commit()
        return count

    def _index_search_text(self, item):
        text = " ".join(
            str(value or "")
            for value in (
                item.get("provider"),
                item.get("brand"),
                item.get("article"),
                self._normalize_article(item.get("article")),
                item.get("name"),
                item.get("warehouse"),
                item.get("logo"),
                item.get("source_brand"),
                item.get("source_code"),
            )
        )
        return self._normalize_search_text(text)

    def _search_groups(self, query_text):
        groups = []
        for part in re.split(r"[,;|]+", str(query_text or "")):
            normalized = self._normalize_search_text(part)
            terms = [
                term for term in normalized.split()
                if len(term) >= 2
            ]
            if terms:
                groups.append(terms)
        if not groups:
            normalized = self._normalize_search_text(query_text)
            terms = [term for term in normalized.split() if len(term) >= 2]
            if terms:
                groups.append(terms)
        return groups

    @staticmethod
    def _normalize_search_text(value):
        text = str(value or "").lower().replace("ё", "е")
        text = re.sub(r"[^a-zа-я0-9]+", " ", text)
        return re.sub(r"\s+", " ", text).strip()

    def check_sources(self, progress_callback=None):
        """Проверяет доступность CSV без поиска условного артикула."""
        available = []
        errors = []
        active_urls = [
            url for url in self.urls
            if (self._profile_for_url(url) or {}).get("enabled", True) is not False
        ]
        total = len(active_urls)
        for index, url in enumerate(active_urls, start=1):
            profile = self._profile_for_url(url) or {}
            if progress_callback:
                progress_callback(index - 1, total, self._source_name(url), None)
            rows = self._load_rows(url, refresh=True)
            if rows is None:
                if self.last_message:
                    errors.append(self.last_message)
                if progress_callback:
                    progress_callback(index, total, self._source_name(url), False)
                continue
            indexed_count = self._replace_source_index(url, rows)
            available.append((url, indexed_count))
            if progress_callback:
                progress_callback(index, total, self._source_name(url), True)
        if available:
            self._sqlite_ready = True
            self.last_message = ""
            total = sum(count for _, count in available)
            return True, f"OK: {len(available)}/{len(self.urls)} CSV, {total} позиций"
        return False, errors[-1] if errors else "Нет доступных CSV"

    def _cache_path(self, url):
        digest = hashlib.sha256(str(url).encode("utf-8")).hexdigest()
        return os.path.join(self.cache_dir, f"{digest}.csv")

    def _load_rows(self, url, refresh=False):
        if str(url).lower().startswith("ftp://") and str(url).endswith("/"):
            return self._load_ftp_directory(url, refresh=refresh)
        cache_path = self._cache_path(url)
        cache_is_fresh = (
            os.path.exists(cache_path)
            and (time.time() - os.path.getmtime(cache_path)) < self.cache_hours * 3600
        )
        if not refresh and cache_is_fresh:
            try:
                with open(cache_path, "rb") as file:
                    content = file.read()
                if urlparse(str(url)).path.lower().endswith(".xlsx"):
                    return self._parse_xlsx(content)
                return self._parse_csv(self._decode_bytes(content), self._profile_for_url(url))
            except FileNotFoundError:
                self.last_message = f"{self.name}: локальный CSV не найден. Нажмите «Проверить Freno CSV» в настройках"
                return None
            except (OSError, UnicodeDecodeError) as exc:
                self.last_message = f"{self.name}: не удалось прочитать локальный CSV: {exc}"
                return None
        try:
            content = self._download_content(url)
        except Exception as exc:
            self.last_message = f"{self.name}: ошибка загрузки {self._safe_url(url)}: {exc}"
            return None

        if urlparse(str(url)).path.lower().endswith(".xlsx"):
            try:
                rows = self._parse_xlsx(content)
                self._save_xlsx_cache(cache_path, content)
                return rows
            except Exception as exc:
                self.last_message = f"{self.name}: ошибка Excel {self._safe_url(url)}: {exc}"
                return None

        try:
            content = self._unpack_content(content)
        except (OSError, zipfile.BadZipFile) as exc:
            self.last_message = f"{self.name}: ошибка архива {self._safe_url(url)}: {exc}"
            return None
        text = self._decode_bytes(content).strip()
        if not text:
            self.last_message = f"{self.name}: пустой CSV для {self._safe_url(url)}"
            return None
        try:
            os.makedirs(self.cache_dir, exist_ok=True)
            temp_path = f"{cache_path}.tmp"
            with open(temp_path, "wb") as file:
                file.write(content or text.encode("utf-8"))
            os.replace(temp_path, cache_path)
        except OSError as exc:
            self.last_message = f"{self.name}: CSV загружен, но не сохранён: {exc}"
            return None
        return self._parse_csv(text, self._profile_for_url(url))

    @staticmethod
    def _save_xlsx_cache(cache_path, content):
        os.makedirs(os.path.dirname(cache_path), exist_ok=True)
        temp_path = f"{cache_path}.tmp"
        with open(temp_path, "wb") as file:
            file.write(content)
        os.replace(temp_path, cache_path)

    @staticmethod
    def _parse_xlsx(content):
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise OSError("для Excel-прайсов установите зависимость openpyxl") from exc
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        sheet = workbook.active
        iterator = sheet.iter_rows(values_only=True)
        headers = [str(value or "").strip() for value in next(iterator, ())]
        rows = []
        for values in iterator:
            row = {
                header: ("" if value is None else value)
                for header, value in zip(headers, values)
                if header
            }
            if any(str(value).strip() for value in row.values()):
                rows.append(row)
        workbook.close()
        return rows

    def _download_content(self, url):
        if str(url).lower().startswith("ftp://"):
            parsed = urlparse(str(url))
            buffer = io.BytesIO()
            encoding = "cp1251" if (parsed.hostname or "").lower() == "ftp.favorit-auto.ru" else "utf-8"
            with ftplib.FTP(timeout=self.timeout, encoding=encoding) as ftp:
                ftp.connect(parsed.hostname, parsed.port or 21)
                ftp.login(unquote(parsed.username or "anonymous"), unquote(parsed.password or ""))
                ftp.retrbinary(f"RETR {unquote(parsed.path)}", buffer.write)
            return buffer.getvalue()
        response = requests.get(
            url,
            timeout=self.timeout,
            headers={"User-Agent": "price_parcer/1.0"},
            proxies={"http": None, "https": None},
        )
        response.raise_for_status()
        return response.content or b""

    def _load_ftp_directory(self, url, refresh=False):
        parsed = urlparse(str(url))
        base_path = unquote(parsed.path or "/")
        rows = []
        try:
            encoding = "cp1251" if (parsed.hostname or "").lower() == "ftp.favorit-auto.ru" else "utf-8"
            with ftplib.FTP(timeout=self.timeout, encoding=encoding) as ftp:
                ftp.connect(parsed.hostname, parsed.port or 21)
                ftp.login(unquote(parsed.username or "anonymous"), unquote(parsed.password or ""))
                ftp.cwd(base_path)
                filenames = ftp.nlst()
        except Exception as exc:
            self.last_message = f"{self.name}: ошибка FTP-каталога: {exc}"
            return None
        allowed = (".csv", ".txt", ".zip", ".gz", ".gzip", ".7z")
        for filename in filenames:
            if not str(filename).lower().endswith(allowed):
                continue
            child_path = f"{base_path.rstrip('/')}/{filename}"
            child_url = parsed._replace(path=child_path).geturl()
            child_rows = self._load_rows(child_url, refresh=refresh)
            if child_rows is None:
                continue
            source_name = humanize_price_name(os.path.basename(str(filename)).rsplit(".", 1)[0], force_prefix=True)
            for row in child_rows:
                row["__source_name"] = source_name
                rows.append(row)
        if not rows:
            self.last_message = f"{self.name}: в FTP-каталоге нет поддерживаемых прайсов"
            return None
        return rows

    @staticmethod
    def _decode_response(response):
        """CSV Freno приходит в UTF-8 с BOM, но без корректной HTTP-кодировки."""
        content = getattr(response, "content", None)
        if content:
            try:
                return content.decode("utf-8-sig")
            except UnicodeDecodeError:
                pass
        return getattr(response, "text", "") or ""

    @staticmethod
    def _decode_bytes(content):
        for encoding in ("utf-8-sig", "cp1251", "utf-16", "latin1"):
            try:
                return bytes(content or b"").decode(encoding)
            except UnicodeDecodeError:
                continue
        return bytes(content or b"").decode("utf-8", errors="replace")

    @staticmethod
    def _unpack_content(content):
        content = bytes(content or b"")
        if content.startswith(b"PK\x03\x04"):
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                names = [
                    name for name in archive.namelist()
                    if not name.endswith("/") and name.lower().endswith((".csv", ".txt"))
                ]
                if not names:
                    names = [name for name in archive.namelist() if not name.endswith("/")]
                if not names:
                    return b""
                return archive.read(names[0])
        if content.startswith(b"\x1f\x8b"):
            return gzip.decompress(content)
        if content.startswith(b"7z\xbc\xaf'\x1c"):
            try:
                import py7zr
            except ImportError as exc:
                raise OSError("для архива 7z установите зависимость py7zr") from exc
            with tempfile.TemporaryDirectory() as directory:
                with py7zr.SevenZipFile(io.BytesIO(content), mode="r") as archive:
                    archive.extractall(path=directory)
                candidates = []
                for root, _, files in os.walk(directory):
                    for filename in files:
                        path = os.path.join(root, filename)
                        if filename.lower().endswith((".csv", ".txt")):
                            candidates.insert(0, path)
                        else:
                            candidates.append(path)
                if not candidates:
                    return b""
                with open(candidates[0], "rb") as file:
                    return file.read()
        return content

    def _parse_csv(self, text, profile=None):
        profile = profile or {}
        has_header = profile.get("has_header", "auto")
        if has_header == "auto":
            first_line = next((line for line in text.splitlines() if line.strip()), "")
            normalized = self._normalize_key(first_line)
            has_header = any(word in normalized for word in (
                "article", "артикул", "brand", "бренд", "price", "цена", "quantity", "остаток"
            ))
        delimiter = self._detect_delimiter(text)
        if has_header:
            reader = csv.DictReader(io.StringIO(text, newline=""), delimiter=delimiter)
        else:
            reader = (
                {str(index): value for index, value in enumerate(row)}
                for row in csv.reader(io.StringIO(text, newline=""), delimiter=delimiter)
            )

        rows = []
        for row in reader:
            if row is None:
                continue
            rows.append({str(k): (v if v is not None else "") for k, v in row.items() if k is not None})
        return rows

    def _detect_delimiter(self, text):
        sample = self._csv_sample(text)
        scored = [
            (self._delimiter_score(sample, delimiter), delimiter)
            for delimiter in (",", ";", "|", "\t")
        ]
        scored.sort(reverse=True)
        best_score, best_delimiter = scored[0]
        if best_score[0] > 0:
            return best_delimiter
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;|\t")
            if dialect.delimiter in {",", ";", "|", "\t"}:
                return dialect.delimiter
        except Exception:
            pass
        return "\t" if "\t" in sample else ";"

    @staticmethod
    def _csv_sample(text, max_lines=250, max_chars=128 * 1024):
        lines = []
        size = 0
        for line in str(text or "").splitlines():
            if not line.strip():
                continue
            lines.append(line)
            size += len(line) + 1
            if len(lines) >= max_lines or size >= max_chars:
                break
        if lines:
            return "\n".join(lines)
        return str(text or "")[:max_chars]

    @staticmethod
    def _delimiter_score(sample, delimiter):
        try:
            widths = [
                len(row)
                for row in csv.reader(io.StringIO(sample, newline=""), delimiter=delimiter)
                if any(str(value or "").strip() for value in row)
            ]
        except Exception:
            return (0, 0, 0, 0)
        useful = [width for width in widths if width > 1]
        if not useful:
            return (0, 0, 0, 0)
        counter = Counter(useful)
        mode_width, mode_count = max(counter.items(), key=lambda item: (item[1], item[0]))
        delimiter_count = str(sample or "").count(delimiter)
        return (mode_count, mode_width, len(useful), delimiter_count)

    def _normalize_row(self, row, source_url):
        if not isinstance(row, dict):
            return None

        self._current_source_url = source_url
        profile = self._profile_for_url(source_url)
        article = self._get_mapped_value(row, "article", [
            "article", "art", "articul", "partnumber", "code", "number", "articlecode", "code1", "артикул"
        ])
        if not article:
            return None

        brand = self._get_mapped_value(row, "brand", [
            "brand", "manuf", "manufacturer", "producer", "brandname", "производитель", "бренд"
        ])
        name = self._get_mapped_value(row, "name", [
            "name", "title", "description", "itemname", "productname", "номенклатура", "наименование"
        ])
        price = self._to_float(self._get_mapped_value(row, "price", ["price", "cost", "price_rub", "price_rur", "price_rus", "value", "цена"]))
        warehouse = self._get_mapped_value(row, "warehouse", [
            "warehouse", "warehouse_name", "stock", "storage", "supplier", "place", "warehouse_name"
        ])
        delivery_value = self._get_mapped_value(row, "days", ["days", "delivery", "delivery_days", "srok", "days_to_delivery"])
        delivery_hours = self._parse_delivery_hours(delivery_value)
        if delivery_hours is None:
            delivery_hours = self._parse_delivery_hours((profile or {}).get("default_days", self.default_days))
        if delivery_hours is None:
            delivery_hours = 0
        days = delivery_hours // 24
        if delivery_hours and delivery_hours % 24:
            days += 1
        quantity = self._get_mapped_value(row, "quantity", ["quantity", "qty", "stockqty", "count", "available", "остаток"])
        multiplicity = self._to_int(self._get_mapped_value(row, "multiplicity", ["multiplicity", "minqty", "min_count"])) or 1
        is_cross = self._is_cross(row)

        return {
            "provider": self._display_source_name(
                row.get("__source_name") or (profile or {}).get("name") or self._source_name(source_url),
                force_prefix=bool(row.get("__source_name")),
            ),
            "brand": str(brand or ""),
            "article": str(article or ""),
            "price": price,
            "days": days,
            "delivery_total_hours": delivery_hours,
            "quantity": str(quantity or "0"),
            "logo": str(warehouse or "-"),
            "warehouse": str(warehouse or ""),
            "name": str(name or ""),
            "multiplicity": multiplicity,
            "is_cross": is_cross,
            "source_url": source_url,
            "source_brand": str(brand or ""),
            "source_code": str(article or ""),
        }

    def _profile_for_url(self, source_url):
        if not source_url:
            return None
        profile = self.url_profiles.get(source_url)
        if profile:
            return profile
        for url, item in (self.url_profiles or {}).items():
            if url and str(url).strip() == str(source_url).strip():
                return item
        return None

    def _parse_delivery_hours(self, value):
        text = str(value or "").strip()
        if not text:
            return None
        if "," in text:
            parts = [part.strip() for part in text.split(",", 1)]
            if len(parts) == 2 and all(re.fullmatch(r"-?\d+", part or "0") for part in parts):
                hours = int(parts[0] or 0)
                days = int(parts[1] or 0)
                return max(0, days * 24 + hours)
            return None
        normalized = text.lower().replace("ё", "е")
        numbers = [int(match) for match in re.findall(r"\d+", normalized)]
        if numbers:
            value = max(numbers)
            if any(token in normalized for token in ("ч", "hour", "час")):
                return max(0, value)
            if any(token in normalized for token in ("д", "day", "сут")):
                return max(0, value * 24)
            if re.fullmatch(r"\s*\d+\s*[-–]\s*\d+\s*", normalized):
                return max(0, value * 24)
        try:
            # Legacy saved CSV settings were stored as days; keep them readable
            # so existing profiles do not become zero after the format change.
            return max(0, int(float(text.replace(",", ".")) * 24))
        except ValueError:
            return None

    def _get_mapped_value(self, row, field_name, aliases):
        if not isinstance(row, dict):
            return ""
        profile = self._profile_for_url(getattr(self, "_current_source_url", None))
        mapped = None
        if profile:
            mapped = profile.get("column_map", {}).get(field_name) or profile.get("column_map", {}).get(self._normalize_key(field_name))
        if not mapped:
            mapped = self.column_map.get(field_name) or self.column_map.get(self._normalize_key(field_name))
        if mapped:
            if mapped in row:
                return row[mapped]
            normalized = {self._normalize_key(k): v for k, v in row.items() if k is not None}
            if self._normalize_key(mapped) in normalized:
                return normalized[self._normalize_key(mapped)]
        if row and all(str(key).isdigit() for key in row):
            positional = {
                "article": "0",
                "brand": "1",
                "name": "2",
                "quantity": "3",
                "price": "4",
                "multiplicity": "5",
                "days": "6",
                "warehouse": "7",
            }
            index = positional.get(field_name)
            if index is not None:
                return row.get(index, "")
        return self._find_value(row, aliases)

    def _source_name(self, source_url):
        profile = self._profile_for_url(source_url) or {}
        if profile.get("name"):
            return self._display_source_name(profile["name"])
        path = urlparse(str(source_url or "")).path.rstrip("/")
        filename = os.path.basename(path)
        if filename:
            stem = filename.rsplit(".", 1)[0]
            if stem and len(stem) <= 80:
                return humanize_price_name(stem, force_prefix=True)
        return self.name

    @staticmethod
    def _display_source_name(name, force_prefix=False):
        return humanize_price_name(name, force_prefix=force_prefix)

    @staticmethod
    def _safe_url(url):
        parsed = urlparse(str(url or ""))
        if parsed.username or parsed.password:
            host = parsed.hostname or ""
            if parsed.port:
                host = f"{host}:{parsed.port}"
            return parsed._replace(netloc=host).geturl()
        return str(url or "")

    def _find_value(self, row, names):
        if not isinstance(row, dict):
            return ""
        normalized = {}
        for key, value in row.items():
            if key is None:
                continue
            normalized[self._normalize_key(key)] = value
        for name in names:
            value = normalized.get(self._normalize_key(name))
            if value not in (None, ""):
                return value
        return ""

    def _normalize_key(self, value):
        return re.sub(r"[^a-zа-я0-9]+", "", str(value or "").lower())

    def _normalize_article(self, value):
        return re.sub(r"[^a-z0-9]+", "", str(value or "").lower())

    def _same_article(self, row_article, target_article):
        if not row_article or not target_article:
            return False
        return self._normalize_article(row_article) == target_article

    def _is_cross(self, row):
        values = []
        for name in ["is_cross", "isanalog", "analog", "cross", "crosscode", "iscross", "crossitem"]:
            value = self._find_value(row, [name])
            if value not in (None, ""):
                values.append(str(value).strip().lower())
        if not values:
            return False
        for value in values:
            if value in {"1", "true", "yes", "y", "да", "дада", "analog", "cross", "кросс", "аналог"}:
                return True
        return False

    def _to_float(self, value, default=0.0):
        try:
            return float(str(value).replace(" ", "").replace(",", "."))
        except Exception:
            return default

    def _to_int(self, value, default=0):
        try:
            return int(float(str(value).replace(" ", "").replace(",", ".")))
        except Exception:
            return default
