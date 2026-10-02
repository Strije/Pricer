import csv
import io
import os
import re
from dataclasses import asdict, dataclass
from decimal import Decimal, InvalidOperation, ROUND_CEILING


@dataclass(frozen=True)
class SalesReportRow:
    row_no: int
    brand: str
    article: str
    name: str
    quantity: int
    article_key: str
    brand_key: str
    source_rows: tuple[int, ...]

    def to_dict(self):
        data = asdict(self)
        data["source_rows"] = list(self.source_rows)
        return data


HEADER_ALIASES = {
    "brand": (
        "brand",
        "бренд",
        "производитель",
        "изготовитель",
        "марка",
        "manuf",
        "manufacturer",
    ),
    "article": (
        "article",
        "артикул",
        "код",
        "кодтовара",
        "номенклатуракод",
        "номер",
        "partnumber",
        "part",
        "sku",
    ),
    "name": (
        "name",
        "название",
        "наименование",
        "товар",
        "номенклатура",
        "description",
        "descr",
    ),
    "quantity": (
        "qty",
        "quantity",
        "колво",
        "количество",
        "кво",
        "шт",
        "продано",
    ),
}


def read_sales_report(path):
    ext = os.path.splitext(str(path or ""))[1].lower()
    if ext in (".xlsx", ".xlsm"):
        rows = _read_xlsx(path)
    else:
        rows = _read_csv(path)
    parsed, errors = parse_sales_report_rows(rows)
    return merge_sales_report_rows(parsed), errors


def parse_sales_report_rows(rows):
    rows = [list(row or []) for row in rows or []]
    rows = [row for row in rows if any(str(cell or "").strip() for cell in row)]
    if not rows:
        return [], ["Файл пустой"]

    header_index, mapping = _detect_header(rows)
    start_index = header_index + 1 if header_index is not None else 0
    if not mapping:
        mapping = {"brand": 0, "article": 1, "name": 2, "quantity": 3}

    parsed = []
    errors = []
    for index, row in enumerate(rows[start_index:], start=start_index + 1):
        brand = _cell(row, mapping.get("brand")).strip()
        article = _cell(row, mapping.get("article")).strip()
        name = _cell(row, mapping.get("name")).strip()
        quantity_text = _cell(row, mapping.get("quantity")).strip()

        if not article:
            errors.append(f"Строка {index}: не указан артикул")
            continue
        quantity = parse_quantity(quantity_text)
        if quantity <= 0:
            errors.append(f"Строка {index}: неверное количество")
            continue
        article_key = normalize_article(article)
        if not article_key:
            errors.append(f"Строка {index}: артикул не содержит букв или цифр")
            continue
        parsed.append(
            SalesReportRow(
                row_no=index,
                brand=brand,
                article=article,
                name=name,
                quantity=quantity,
                article_key=article_key,
                brand_key=normalize_brand(brand),
                source_rows=(index,),
            )
        )
    return parsed, errors


def merge_sales_report_rows(rows):
    merged = {}
    order = []
    for row in rows or []:
        key = (row.brand_key, row.article_key)
        if key not in merged:
            merged[key] = row
            order.append(key)
            continue
        current = merged[key]
        merged[key] = SalesReportRow(
            row_no=current.row_no,
            brand=current.brand or row.brand,
            article=current.article,
            name=current.name or row.name,
            quantity=current.quantity + row.quantity,
            article_key=current.article_key,
            brand_key=current.brand_key,
            source_rows=current.source_rows + row.source_rows,
        )
    return [merged[key] for key in order]


def parse_quantity(value):
    text = str(value or "").strip()
    if not text:
        return 0
    match = re.search(r"\d+(?:[.,]\d+)?", text)
    if not match:
        return 0
    try:
        number = Decimal(match.group(0).replace(",", "."))
    except (InvalidOperation, ValueError):
        return 0
    if number <= 0:
        return 0
    return int(number.to_integral_value(rounding=ROUND_CEILING))


def normalize_article(value):
    return re.sub(r"[^A-ZА-Я0-9]+", "", str(value or "").upper().replace("Ё", "Е"))


def normalize_brand(value):
    return re.sub(r"[^A-ZА-Я0-9]+", "", str(value or "").upper().replace("Ё", "Е"))


def _read_csv(path):
    content = None
    with open(path, "rb") as file:
        raw = file.read()
    for encoding in ("utf-8-sig", "cp1251", "utf-16"):
        try:
            content = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if content is None:
        content = raw.decode("utf-8", errors="replace")
    delimiter = _detect_delimiter(content)
    return list(csv.reader(io.StringIO(content, newline=""), delimiter=delimiter))


def _read_xlsx(path):
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise OSError("для Excel-файлов нужна зависимость openpyxl") from exc
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    rows = []
    for row in sheet.iter_rows(values_only=True):
        rows.append(["" if value is None else str(value) for value in row])
    workbook.close()
    return rows


def _detect_delimiter(text):
    sample = "\n".join(str(text or "").splitlines()[:80])
    delimiters = [",", ";", "\t", "|"]
    scores = []
    for delimiter in delimiters:
        counts = [line.count(delimiter) for line in sample.splitlines() if line.strip()]
        positive = [count for count in counts if count > 0]
        if not positive:
            scores.append((0, 0, delimiter))
            continue
        common = max(set(positive), key=positive.count)
        scores.append((positive.count(common), common, delimiter))
    scores.sort(reverse=True)
    return scores[0][2] if scores and scores[0][0] else ";"


def _detect_header(rows):
    best_index = None
    best_mapping = {}
    for index, row in enumerate(rows[:10]):
        mapping = {}
        normalized_headers = [_header_key(cell) for cell in row]
        for field, aliases in HEADER_ALIASES.items():
            alias_keys = {_header_key(alias) for alias in aliases}
            for col, header in enumerate(normalized_headers):
                if _header_matches(header, alias_keys):
                    mapping[field] = col
                    break
        score = sum(1 for field in ("article", "quantity") if field in mapping)
        score += sum(1 for field in ("brand", "name") if field in mapping)
        if score > len(best_mapping):
            best_index = index
            best_mapping = mapping
    if "article" not in best_mapping or "quantity" not in best_mapping:
        return None, {}
    return best_index, best_mapping


def _header_key(value):
    return re.sub(r"[^a-zа-я0-9]+", "", str(value or "").lower().replace("ё", "е"))


def _header_matches(header, alias_keys):
    if header in alias_keys:
        return True
    for alias in alias_keys:
        if len(alias) >= 5 and (header.startswith(alias) or alias in header):
            return True
    return False


def _cell(row, index):
    if index is None:
        return ""
    try:
        return str(row[int(index)] or "")
    except (IndexError, TypeError, ValueError):
        return ""
