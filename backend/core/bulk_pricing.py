"""Подготовка произвольного списка к массовой проценке.

Модуль намеренно не тянет Qt и сеть: здесь только чтение файла, приведение
колонок и наименований к рабочему виду и безопасное сохранение результата.
Сама проценка живёт в bulk_price_cli.py.
"""
import csv
import html
import io
import os
import re

from sales_report_importer import (
    _detect_header,
    normalize_article,
    normalize_brand,
    parse_quantity,
)

RESULT_COLUMNS = (
    "Цена поиска",
    "Закуп",
    "Поставщик",
    "Срок, дн",
    "Наличие",
    "Статус",
)
SEARCH_COLUMNS = ("Бренд", "Артикул", "Наименование", "Кол-во")

# Выгрузки 1С обрезают длинные названия брендов, и поставщик такой бренд не узнаёт.
BRAND_FIXUPS = {
    "GENERAL": ("GENERAL MOTORS", "GM"),
    "MERCEDES": ("MERCEDES-BENZ",),
    "LAND": ("LAND ROVER",),
    "ALFA": ("ALFA ROMEO",),
}

# Латиница, которой в 1С часто «портят» русские слова: PОЛИКОПОДШИПНИК и т.п.
_HOMOGLYPHS = str.maketrans(
    "ACEHKMOPTXBacepxy",
    "АСЕНКМОРТХВасерху",  # noqa: RUF001 - намеренная замена латиницы на кириллицу
)


def read_any(path, columns=None):
    """Читает csv/xlsx/xls-html и возвращает (строки, заголовок исходника).

    Строка — dict с ключами brand/article/name/quantity/extras.
    `columns` — ручная раскладка вида "brand,article,quantity,price,-,name,cell",
    прочерк пропускает колонку.
    """
    raw_rows = _read_raw(path)
    raw_rows = [row for row in raw_rows if any(str(cell or "").strip() for cell in row)]
    if not raw_rows:
        return [], []

    if columns:
        mapping = _parse_columns_option(columns)
        header_index = None
    else:
        header_index, mapping = _detect_header(raw_rows)
        if not mapping:
            mapping = guess_layout(raw_rows)
    if "article" not in mapping:
        raise ValueError(
            "не удалось определить колонку с артикулом — задайте раскладку через --columns"
        )

    header = list(raw_rows[header_index]) if header_index is not None else []
    start = (header_index + 1) if header_index is not None else 0

    rows = []
    for index, raw in enumerate(raw_rows[start:], start=start + 1):
        article = _cell(raw, mapping.get("article")).strip()
        if not article or not normalize_article(article):
            continue
        quantity = parse_quantity(_cell(raw, mapping.get("quantity"))) or 1
        rows.append(
            {
                "row_no": index,
                "brand": _cell(raw, mapping.get("brand")).strip(),
                "article": article,
                "name": pretty_name(_cell(raw, mapping.get("name"))),
                "quantity": quantity,
                "extras": list(raw),
            }
        )
    return rows, header


def brand_candidates(brand):
    """Бренд как в файле, затем известные расшифровки обрезанных названий."""
    brand = str(brand or "").strip()
    if not brand:
        return []
    out = [brand]
    for fixup in BRAND_FIXUPS.get(brand.upper(), ()):
        if normalize_brand(fixup) != normalize_brand(brand):
            out.append(fixup)
    return out


def pretty_name(name):
    """ВЕРХНИЙ РЕГИСТР -> нормальный. Латинские марки и индексы не трогаем."""
    name = re.sub(r"\s+", " ", str(name or "")).strip()
    if not name:
        return ""
    out = []
    first_word_done = False
    for token in name.split(" "):
        token = fix_homoglyphs(token)
        letters = re.sub(r"[^A-Za-zА-Яа-яЁё]", "", token)
        has_cyr = bool(re.search(r"[А-Яа-яЁё]", token))
        has_lat = bool(re.search(r"[A-Za-z]", token))
        has_digit = bool(re.search(r"\d", token))
        if not has_cyr or has_lat or has_digit or len(letters) <= 2:
            out.append(token)
            if letters:
                first_word_done = True
            continue
        lowered = token.lower()
        if not first_word_done:
            lowered = re.sub(r"([а-яё])", lambda m: m.group(1).upper(), lowered, count=1)
            first_word_done = True
        out.append(lowered)
    return " ".join(out)


def fix_homoglyphs(token):
    """PОЛИКО... -> РОЛИКО..., но LEDO и D3 остаются как есть."""
    if not re.search(r"[А-Яа-яЁё]", token):
        return token
    fixed = token.translate(_HOMOGLYPHS)
    return fixed if not re.search(r"[A-Za-z]", fixed) else token


def safe_save(workbook, path, attempts=5):
    """Excel держит открытый файл заблокированным — тогда пишем в соседний."""
    base, ext = os.path.splitext(path)
    candidates = [path] + [f"{base}_{n}{ext}" for n in range(2, attempts + 1)]
    last_error = None
    for candidate in candidates:
        try:
            workbook.save(candidate)
            return candidate
        except PermissionError as exc:
            last_error = exc
    raise PermissionError(
        f"файл занят, свободного имени не нашлось: {path} ({last_error})"
    )


def guess_layout(rows):
    """Раскладка для файла без заголовка: по содержимому колонок."""
    sample = rows[: min(len(rows), 200)]
    width = max(len(row) for row in sample)
    stats = [_column_stats(sample, col) for col in range(width)]

    mapping = {}
    text_cols = [s for s in stats if s["wordy"] >= 0.5]
    if text_cols:
        name = max(text_cols, key=lambda s: s["avg_len"])
        mapping["name"] = name["col"]

    code_cols = [
        s for s in stats
        if s["col"] != mapping.get("name")
        and s["code"] >= 0.7
        and s["avg_len"] >= 3
        # колонки с копейками — это цены, а не артикулы
        and s["decimals"] < 0.5
    ]
    if code_cols:
        # Артикул отличает плотность цифр: 0010098SX — это 0.78, а склад-
        # ская ячейка ГС-2-2-К1 всего 0.33 при той же длине. Бренд обычно левее.
        article = max(code_cols, key=lambda s: (s["digit_density"], s["avg_len"]))
        mapping["article"] = article["col"]
        rest = [s for s in code_cols if s["col"] != article["col"]]
        if rest:
            mapping["brand"] = min(rest, key=lambda s: s["col"])["col"]

    int_cols = [
        s for s in stats
        if s["col"] not in mapping.values() and s["ints"] >= 0.9 and s["avg_value"] <= 1000
    ]
    if int_cols:
        mapping["quantity"] = min(int_cols, key=lambda s: s["avg_value"])["col"]
    return mapping


def _column_stats(rows, col):
    values = [str(_cell(row, col) or "").strip() for row in rows]
    values = [value for value in values if value]
    total = max(1, len(values))
    numbers = []
    for value in values:
        try:
            numbers.append(float(value.replace(" ", "").replace(",", ".")))
        except ValueError:
            pass
    return {
        "col": col,
        "avg_len": sum(len(value) for value in values) / total,
        # «словесная» колонка: буквы и пробелы, как в наименовании
        "wordy": sum(1 for value in values if " " in value and re.search(r"[A-Za-zА-Яа-яЁё]", value)) / total,
        # «кодовая»: без пробелов, есть буквы или цифры
        "code": sum(1 for value in values if " " not in value and re.search(r"[A-Za-z0-9А-Яа-я]", value)) / total,
        # доля цифр внутри значения, а не «есть ли цифры вообще»
        "digit_density": sum(
            len(re.findall(r"\d", value)) / len(value) for value in values
        ) / total,
        "decimals": sum(
            1 for value in values if re.fullmatch(r"-?\d+[.,]\d+", value.replace(" ", ""))
        ) / total,
        "ints": sum(1 for value in values if re.fullmatch(r"\d+", value)) / total,
        "avg_value": (sum(numbers) / len(numbers)) if numbers else 10**9,
    }


def _parse_columns_option(columns):
    known = {"brand", "article", "name", "quantity", "qty"}
    mapping = {}
    for index, token in enumerate(str(columns).split(",")):
        token = token.strip().lower()
        if token in ("", "-", "skip"):
            continue
        if token == "qty":
            token = "quantity"
        if token not in known:
            continue
        mapping.setdefault(token, index)
    return mapping


def _read_raw(path):
    ext = os.path.splitext(str(path or ""))[1].lower()
    with open(path, "rb") as file:
        head = file.read(512)
    # выгрузки 1С зовутся .xls, а внутри HTML-таблица
    if head.lstrip()[:1] == b"<":
        return _read_html_table(path)
    if ext in (".xlsx", ".xlsm"):
        return _read_xlsx(path)
    return _read_csv(path)


def _read_html_table(path):
    raw = open(path, "rb").read()
    encoding = "cp1251" if re.search(rb"windows-1251", raw[:400], re.I) else "utf-8"
    text = raw.decode(encoding, errors="replace")
    rows = []
    for tr in re.findall(r"<tr[^>]*>(.*?)</tr>", text, re.S | re.I):
        cells = re.findall(r"<t[dh][^>]*>(.*?)</t[dh]>", tr, re.S | re.I)
        cleaned = []
        for cell in cells:
            cell = re.sub(r"<[^>]+>", " ", cell)
            cell = html.unescape(cell).replace("\xa0", " ")
            cleaned.append(re.sub(r"\s+", " ", cell).strip())
        if cleaned:
            rows.append(cleaned)
    return rows


def _read_xlsx(path):
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    rows = [
        ["" if value is None else str(value) for value in row]
        for row in sheet.iter_rows(values_only=True)
    ]
    workbook.close()
    return rows


def _read_csv(path):
    raw = open(path, "rb").read()
    content = None
    for encoding in ("utf-8-sig", "cp1251", "utf-16"):
        try:
            content = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if content is None:
        content = raw.decode("utf-8", errors="replace")
    delimiter = ";"
    sample = "\n".join(content.splitlines()[:50])
    best = 0
    for option in (",", ";", "\t", "|"):
        count = sample.count(option)
        if count > best:
            best, delimiter = count, option
    return list(csv.reader(io.StringIO(content, newline=""), delimiter=delimiter))


def _cell(row, index):
    if index is None:
        return ""
    try:
        return str(row[int(index)] or "")
    except (IndexError, TypeError, ValueError):
        return ""
