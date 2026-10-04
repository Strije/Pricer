"""Справочник синонимов брендов ABCP: бренд и все его написания («MANN-FILTER», «MANN FILTER», «MANN+HUMMEL»).

Десктоп носит его с собой (resources/brand_reference_abcp.json, ~19 тыс. брендов, ~12 тыс. синонимов),
и по нему резолвер (core/brand_aliases.py) сводит разные написания у поставщиков к одному бренду.
Без него веб узнаёт как один бренд лишь ~10% пар «бренд — другое написание» (проверено на выгрузке
ABCP 04.10.2026), отсюда хуже голосование за бренд и сверка «это искомая деталь».

Источник — выгрузка справочника брендов из админки ABCP (.xls/.xlsx, лист «Бренды»: «Бренд»,
«Алиас(Синоним)» через запятую) или JSON десктопа. Файл ставится в папку данных сервера
(PRICER_DATA_DIR, по умолчанию var/brands) командой `python -m app.manage import-brands <файл>`
(на сервере — `setup.sh import-brands <файл>`), а не в репозиторий: это данные ABCP.
Лист «Группы брендов» (OPEL + ACDelco, TOYOTA + DAIHATSU) не переносится: это разные
производители, слияние их в один бренд путало бы искомую деталь с аналогом.
"""
import json
import os

FILE_NAME = "brand_reference_abcp.json"


def _split_aliases(value):
    return [part.strip() for part in str(value or "").split(",") if part.strip()]


def _rows_to_items(rows):
    items, seen = [], set()
    for row in rows:
        cells = list(row) + ["", ""]
        name = str(cells[0] if cells[0] is not None else "").strip()
        if not name or name in ("Бренд", "Brand"):  # строка заголовка
            continue
        key = name.upper()
        if key in seen:
            continue
        seen.add(key)
        items.append({"name": name, "aliases": _split_aliases(cells[1])})
    return items


def _pick_sheet(names):
    for name in names:
        if str(name).strip().lower() == "бренды":
            return name
    return names[0] if names else None


def read_reference(path):
    """Файл справочника -> [{"name", "aliases"}]."""
    ext = os.path.splitext(path)[1].lower()
    if ext == ".json":
        with open(path, encoding="utf-8-sig") as file:
            data = json.load(file)
        if not isinstance(data, list):
            raise ValueError("в JSON ожидается список брендов")
        rows = [(item.get("name"), ", ".join(item.get("aliases") or []) if isinstance(item.get("aliases"), list)
                 else item.get("aliases")) for item in data if isinstance(item, dict)]
        return _rows_to_items(rows)
    if ext == ".xlsx":
        from openpyxl import load_workbook

        book = load_workbook(path, read_only=True, data_only=True)
        try:
            sheet = book[_pick_sheet(book.sheetnames)]
            return _rows_to_items(sheet.iter_rows(values_only=True))
        finally:
            book.close()
    if ext == ".xls":
        try:
            import xlrd
        except ImportError:
            raise ValueError("для .xls нужен пакет xlrd (pip install xlrd) — или сохраните файл как .xlsx")
        # выгрузка ABCP формально повреждена (xlrd: «Workbook corruption»), при этом данные целы
        with open(os.devnull, "w") as quiet:
            book = xlrd.open_workbook(path, ignore_workbook_corruption=True, logfile=quiet)
        sheet = book.sheet_by_name(_pick_sheet(book.sheet_names()))
        return _rows_to_items(sheet.row_values(r) for r in range(sheet.nrows))
    raise ValueError("поддерживаются .xls, .xlsx и .json")


def install(path, brand_dir):
    """Ставит справочник в папку данных. Возвращает сводку: сколько брендов и синонимов, что изменилось."""
    items = read_reference(path)
    if len(items) < 100:
        raise ValueError(f"в файле всего {len(items)} брендов — это не похоже на справочник ABCP")
    target = os.path.join(brand_dir, FILE_NAME)
    before = set()
    if os.path.exists(target):
        try:
            with open(target, encoding="utf-8") as file:
                before = {str(item.get("name") or "").upper() for item in json.load(file) if isinstance(item, dict)}
        except (OSError, ValueError):
            before = set()
    os.makedirs(brand_dir, exist_ok=True)
    tmp = target + ".tmp"
    with open(tmp, "w", encoding="utf-8") as file:
        json.dump(items, file, ensure_ascii=False)
    os.replace(tmp, target)  # резолвер не увидит недописанный файл
    names = {item["name"].upper() for item in items}
    return {"brands": len(items), "aliases": sum(len(item["aliases"]) for item in items),
            "added": len(names - before) if before else None, "removed": len(before - names) if before else None,
            "path": target}
