"""Прайс-листы: источники (ссылка, FTP; почта — отдельным шагом), загрузка на сервере по расписанию,
разметка колонок, хранение строк в базе и поиск по ним как по обычному поставщику.

Скачивание, FTP-каталоги, архивы (ZIP, GZ, 7z), CSV/XLSX и угадывание колонок по заголовкам —
из десктопа (core/url_csv_provider.py: _load_rows, _normalize_row), чтобы разбор совпадал.
Новое: строки — в базе с индексом «организация + артикул» вместо файлового кэша, у каждого
источника своё расписание, разметка колонок вручную, отчёт загрузки, устаревание.
"""
import csv
import datetime
import io
import os
import re
import zipfile
from urllib.parse import unquote, urlparse

from sqlalchemy import delete, insert

from app import db
from price_names import humanize_price_name
from url_csv_provider import UrlCsvProvider

# Поля строки прайса (как читает десктоп): (код, подпись, обязательное)
FIELDS = [("article", "Артикул", True), ("brand", "Бренд", False), ("name", "Наименование", False),
          ("price", "Цена", True), ("quantity", "Остаток", False), ("multiplicity", "Кратность", False),
          ("days", "Срок поставки", False), ("warehouse", "Склад", False)]
# Заголовки, которые десктоп узнаёт сам (url_csv_provider._normalize_row)
ALIASES = {
    "article": ["article", "art", "articul", "partnumber", "code", "number", "articlecode", "code1", "артикул"],
    "brand": ["brand", "manuf", "manufacturer", "producer", "brandname", "производитель", "бренд"],
    "name": ["name", "title", "description", "itemname", "productname", "номенклатура", "наименование"],
    "price": ["price", "cost", "price_rub", "price_rur", "price_rus", "value", "цена"],
    "quantity": ["quantity", "qty", "stockqty", "count", "available", "остаток"],
    "multiplicity": ["multiplicity", "minqty", "min_count", "кратность"],
    "days": ["days", "delivery", "delivery_days", "srok", "days_to_delivery", "срок"],
    "warehouse": ["warehouse", "warehouse_name", "stock", "storage", "supplier", "place", "склад"],
}
WARN_DAYS = 2   # прайс не обновлялся — предупреждение
STALE_DAYS = 7  # старше — в выдачу не попадает
INSERT_CHUNK = 2000
PREVIEW_SCAN = 20000  # предпросмотр: дальше строки не считаем
SCHEDULES = [1, 2, 4, 6, 12, 24, 48, 168]


def article_key(value):
    """Как в десктопе (_normalize_article): только латиница и цифры, без регистра."""
    return re.sub(r"[^a-z0-9]+", "", str(value or "").lower())[:80]


def location_hint(location):
    """Что показать вместо адреса: схема, сервер и имя файла — без логина, пароля и токенов."""
    parsed = urlparse(str(location or ""))
    name = os.path.basename(unquote(parsed.path.rstrip("/"))) or "/"
    return f"{parsed.scheme}://{parsed.hostname or '?'}/…/{name}"[:300] if parsed.scheme else "?"


def _provider(location, source_settings, name="Прайс", timeout=90):
    settings = source_settings or {}
    profile = {"name": name, "enabled": True, "has_header": settings.get("has_header", "auto"),
               "column_map": dict(settings.get("column_map") or {}), "default_days": settings.get("default_days", "")}
    return UrlCsvProvider(urls=[location], name=name, timeout=timeout, column_map=profile["column_map"],
                          default_days=settings.get("default_days") or 0, url_profiles={location: profile},
                          cache_hours=1)


def _iter_csv(provider, text, profile):
    """Как UrlCsvProvider._parse_csv, но по одной строке: прайс в сотни тысяч строк не держим списком."""
    has_header = (profile or {}).get("has_header", "auto")
    if has_header == "auto":
        first_line = next((line for line in io.StringIO(text) if line.strip()), "")
        normalized = provider._normalize_key(first_line)
        has_header = any(word in normalized for word in (
            "article", "артикул", "brand", "бренд", "price", "цена", "quantity", "остаток"))
    delimiter = provider._detect_delimiter(text[:1024 * 1024])
    if has_header:
        reader = csv.DictReader(io.StringIO(text, newline=""), delimiter=delimiter)
    else:
        reader = ({str(index): value for index, value in enumerate(row)}
                  for row in csv.reader(io.StringIO(text, newline=""), delimiter=delimiter))
    for row in reader:
        if row is not None:
            yield {str(k): (v if v is not None else "") for k, v in row.items() if k is not None}


def _iter_xlsx(content):
    """Как UrlCsvProvider._parse_xlsx, но по одной строке (openpyxl в режиме только чтения)."""
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("для Excel-прайсов установите зависимость openpyxl") from exc
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        iterator = workbook.active.iter_rows(values_only=True)
        headers = [str(value or "").strip() for value in next(iterator, ())]
        for values in iterator:
            row = {h: ("" if v is None else v) for h, v in zip(headers, values) if h}
            if any(str(v).strip() for v in row.values()):
                yield row
    finally:
        workbook.close()


def iter_content(provider, content, filename, location):
    """Строки скачанного файла (вложение письма, файл по ссылке или из FTP-каталога)."""
    lower = str(filename or "").lower()
    data = bytes(content or b"")
    if lower.endswith(".xls") and not data.startswith(b"PK"):
        raise RuntimeError("старый формат XLS не поддерживается — попросите поставщика присылать XLSX или CSV")
    if lower.endswith(".rar") or data.startswith(b"Rar!"):
        raise RuntimeError("архив RAR не поддерживается — попросите поставщика присылать ZIP")
    if data.startswith(b"PK") and not lower.endswith(".xlsx"):
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            inner = [n for n in archive.namelist() if not n.endswith("/")]
            xlsx = [n for n in inner if n.lower().endswith(".xlsx")]
            if xlsx and not any(n.lower().endswith((".csv", ".txt")) for n in inner):
                data, lower = archive.read(xlsx[0]), xlsx[0].lower()
    if lower.endswith(".xlsx") or (data.startswith(b"PK") and b"xl/" in data[:2000]):
        return _iter_xlsx(data)
    text = provider._decode_bytes(provider._unpack_content(data))
    del data
    if not text or text.isspace():
        raise RuntimeError("файл прайса пустой")
    return _iter_csv(provider, text, provider._profile_for_url(location))


def open_rows(location, source_settings, name="Прайс"):
    """(provider, строки) по адресу: файл по ссылке или FTP, либо все прайсы FTP-каталога (адрес на «/»)."""
    provider = _provider(location, source_settings, name)
    if location.lower().startswith("ftp://") and location.endswith("/"):
        return provider, _iter_ftp_directory(provider, location)
    try:
        content = provider._download_content(location)
    except Exception as exc:
        raise RuntimeError(f"{name}: ошибка загрузки {provider._safe_url(location)}: {exc}") from exc
    return provider, iter_content(provider, content, urlparse(location).path, location)


def _iter_ftp_directory(provider, location):
    import ftplib

    parsed = urlparse(location)
    base = unquote(parsed.path or "/")
    encoding = "cp1251" if (parsed.hostname or "").lower() == "ftp.favorit-auto.ru" else "utf-8"
    try:
        with ftplib.FTP(timeout=provider.timeout, encoding=encoding) as ftp:
            ftp.connect(parsed.hostname, parsed.port or 21)
            ftp.login(unquote(parsed.username or "anonymous"), unquote(parsed.password or ""))
            ftp.cwd(base)
            names = [n for n in ftp.nlst() if str(n).lower().endswith((".csv", ".txt", ".zip", ".gz", ".gzip", ".7z"))]
    except Exception as exc:
        raise RuntimeError(f"{provider.name}: ошибка FTP-каталога: {exc}") from exc
    if not names:
        raise RuntimeError(f"{provider.name}: в FTP-каталоге нет поддерживаемых прайсов")
    for filename in names:  # по одному файлу: скачали, разобрали, отпустили
        child = parsed._replace(path=f"{base.rstrip('/')}/{filename}").geturl()
        try:
            rows = iter_content(provider, provider._download_content(child), filename, location)
        except Exception:
            continue
        source_name = humanize_price_name(os.path.basename(str(filename)).rsplit(".", 1)[0], force_prefix=True)
        for row in rows:
            row["__source_name"] = source_name
            yield row


def guess_mapping(headers):
    """Какая колонка за какое поле — по названиям заголовков, как угадывает десктоп."""
    def norm(value):
        return re.sub(r"[^a-zа-я0-9]+", "", str(value or "").lower())
    out = {}
    for field, aliases in ALIASES.items():
        keys = [norm(a) for a in aliases]
        for header in headers:
            h = norm(header)
            if h and (h in keys or any(h.startswith(k) for k in keys if len(k) >= 4)) and header not in out.values():
                out[field] = header
                break
    return out


def preview(rows, source_settings, limit=20):
    """Первые строки и колонки для разметки; строки считаем до PREVIEW_SCAN, дальше — «больше»."""
    rows, sample, total = iter(rows), [], 0
    for row in rows:
        total += 1
        if len(sample) < 200:
            sample.append(row)
        if total >= PREVIEW_SCAN:
            break
    more = total >= PREVIEW_SCAN and next(rows, None) is not None
    headers = []
    for row in sample:
        for key in row:
            if key not in headers and not str(key).startswith("__"):
                headers.append(key)
    return {"headers": headers, "rows": [[str(row.get(h, ""))[:80] for h in headers] for row in sample[:limit]],
            "total": total, "total_more": more, "guess": guess_mapping(headers),
            "mapping": dict((source_settings or {}).get("column_map") or {})}


def normalize_row(provider, row, location, warehouse=""):
    """Строка файла -> (предложение движка, None) или (None, причина пропуска)."""
    item = provider._normalize_row(row, location)
    if item is None:
        return None, "нет артикула"
    if not item.get("price") or float(item["price"]) <= 0:
        return None, "нет цены"
    if warehouse and not item.get("warehouse"):
        item["warehouse"] = item["logo"] = warehouse
    item.pop("source_url", None)  # адрес с паролем в предложение не кладём
    return item, None


def mailbox_for(session, box, source):
    mail = (source.settings or {}).get("mail") or {}
    account = session.get(db.SupplierAccount, int(mail.get("mailbox_id") or 0)) if mail.get("mailbox_id") else None
    if account is None or account.organization_id != source.organization_id or account.section != "mailbox":
        raise RuntimeError("выберите почтовый ящик (Настройки → Сервисы → Почтовые ящики)")
    return {**(account.config or {}), **box.open(account.secrets_sealed)}, mail


def open_mail(session, box, source, skip_message_id=None):
    """(provider, location, строки, письмо) по правилу источника «почта».
    Если самое новое подходящее письмо — skip_message_id (уже загружено), вложение не скачиваем: строки None."""
    from mail_prices import find_latest

    mailbox, rule = mailbox_for(session, box, source)
    content, filename, message_id, sent = find_latest(mailbox, rule, skip_message_id=skip_message_id)
    letter = {"message_id": message_id, "file": filename or "", "sent": sent}
    if content is None:
        return None, "", None, letter
    location = f"mail://{filename}"
    provider = _provider(location, source.settings, source.name)
    return provider, location, iter_content(provider, content, filename, location), letter


def load(session, box, source, force=False):
    """Загрузка одного источника: скачать (или взять из письма), разобрать и заменить строки в базе — пачками,
    в одной транзакции (поиск до конца загрузки видит старый прайс), с отчётом."""
    started = datetime.datetime.now()
    letter = None
    try:
        if source.kind == "email":
            previous = source.status or {}
            same = previous.get("message_id") if not force and previous.get("ok") and source.loaded_at else None
            provider, location, rows, letter = open_mail(session, box, source, skip_message_id=same)
            if rows is None:
                # то же письмо, что в прошлый раз: строки на месте, свежесть — по дате того письма
                source.status = {**previous, "checked_at": db.utcnow().isoformat(timespec="seconds"),
                                 "message": previous.get("message", "").split(" · новых писем нет")[0] + " · новых писем нет"}
                session.commit()
                return source.status
        else:
            location = box.open(source.location_sealed).get("location", "") if source.location_sealed else ""
            if not location:
                raise RuntimeError("не указан адрес прайса")
            provider, rows = open_rows(location, source.settings, source.name)
        warehouse = (source.settings or {}).get("warehouse", "")
        session.execute(delete(db.PriceRow).where(db.PriceRow.source_id == source.id))
        seen, count, reasons, batch = 0, 0, {}, []
        for row in rows:
            seen += 1
            item, reason = normalize_row(provider, row, location, warehouse)
            if reason:
                reasons[reason] = reasons.get(reason, 0) + 1
                continue
            batch.append({"organization_id": source.organization_id, "source_id": source.id,
                          "article_key": article_key(item.get("article")), "brand": str(item.get("brand") or "")[:100],
                          "item": item})
            count += 1
            if len(batch) >= INSERT_CHUNK:
                session.execute(insert(db.PriceRow), batch)
                batch = []
        if batch:
            session.execute(insert(db.PriceRow), batch)
        if not seen:
            raise RuntimeError("в файле нет строк — прежний прайс оставлен")
        if not count:
            raise RuntimeError("ни одной строки не разобрано — проверьте колонки (артикул и цена); прежний прайс оставлен")
        source.loaded_at = db.utcnow()
        skipped = sum(reasons.values())
        status = {"ok": True, "rows": count, "skipped": skipped, "reasons": reasons,
                  "message": f"загружено {count} строк" + (f", пропущено {skipped}" if reasons else "")}
        if letter:
            status.update(letter)
            status["message"] += f" · письмо {letter['sent'][:16].replace('T', ' ')}, файл {letter['file']}"
    except Exception as exc:
        session.rollback()  # строки прежнего прайса остаются
        source = session.get(db.PriceSource, source.id)
        status = {"ok": False, "message": str(exc)[:300]}
    status["at"] = db.utcnow().isoformat(timespec="seconds")
    status["seconds"] = round((datetime.datetime.now() - started).total_seconds(), 1)
    source.status = status
    session.commit()
    return status


def due(source, now=None):
    if not source.enabled:
        return False
    if source.loaded_at is None:
        return not (source.status or {}).get("at") or _retry_due(source, now)
    return (now or db.utcnow()) - source.loaded_at >= datetime.timedelta(hours=max(1, source.schedule_hours or 24))


def _retry_due(source, now):
    """После ошибки не долбим поставщика каждую минуту: повтор через час."""
    try:
        at = datetime.datetime.fromisoformat(source.status["at"])
    except (KeyError, TypeError, ValueError):
        return True
    return (now or db.utcnow()) - at >= datetime.timedelta(hours=1)


def age_days(source, now=None):
    return None if source.loaded_at is None else ((now or db.utcnow()) - source.loaded_at).total_seconds() / 86400


def view(source):
    age = age_days(source)
    state = ("off" if not source.enabled else "bad" if (source.status or {}).get("ok") is False and source.loaded_at is None
             else "stale" if age is not None and age > STALE_DAYS else "warn" if age is not None and age > WARN_DAYS
             else "ok" if source.loaded_at else "new")
    return {"id": source.id, "name": source.name, "kind": source.kind, "enabled": source.enabled,
            "location_hint": source.location_hint, "has_location": bool(source.location_sealed),
            "settings": source.settings or {}, "schedule_hours": source.schedule_hours, "status": source.status or {},
            "loaded_at": source.loaded_at.isoformat(timespec="minutes") if source.loaded_at else None,
            "age_days": round(age, 1) if age is not None else None, "state": state}


def import_desktop(session, box, organization_id, settings):
    """Источники из раздела url_csv десктопа (urls + url_profiles): названия, сроки, вкл./выкл."""
    urls = [str(u).strip() for u in (settings.get("urls") or []) if str(u).strip()]
    profiles = settings.get("url_profiles") or {}
    existing = {box.open(s.location_sealed).get("location") for s in session.query(db.PriceSource)
                .filter_by(organization_id=organization_id) if s.location_sealed}
    added = 0
    for url in urls:
        if url in existing:
            continue
        profile = profiles.get(url) or {}
        name = str(profile.get("name") or "").strip() or os.path.basename(urlparse(url).path).rsplit(".", 1)[0] or "Прайс"
        session.add(db.PriceSource(
            organization_id=organization_id, name=name[:200], kind="ftp" if url.lower().startswith("ftp") else "url",
            enabled=profile.get("enabled", True) is not False, location_sealed=box.seal({"location": url}),
            location_hint=location_hint(url), schedule_hours=max(1, int(settings.get("cache_hours") or 24)),
            settings={"has_header": profile.get("has_header", "auto"), "column_map": profile.get("column_map") or {},
                      "default_days": profile.get("default_days") or settings.get("default_days") or ""}))
        added += 1
    session.commit()
    return added


class PriceDbProvider(UrlCsvProvider):
    """Поставщик движка «Прайс-листы»: ищет в загруженных прайсах организации (вместо файлов десктопа).
    Подкласс UrlCsvProvider — чтобы поиск аналогов по локальным прайсам (_search_crosses) шёл и сюда."""

    DISPLAY_NAME = "Прайс-листы"

    def __init__(self, sessionmaker, organization_id):
        super().__init__(urls=["db://prices"], name="Прайс-листы")
        self.sessionmaker, self.organization_id = sessionmaker, organization_id

    def _rows(self, keys):
        fresh = db.utcnow() - datetime.timedelta(days=STALE_DAYS)
        with self.sessionmaker() as session:
            return (session.query(db.PriceRow.article_key, db.PriceRow.item)
                    .join(db.PriceSource, db.PriceSource.id == db.PriceRow.source_id)
                    .filter(db.PriceRow.organization_id == self.organization_id, db.PriceRow.article_key.in_(list(keys)),
                            db.PriceSource.enabled.is_(True), db.PriceSource.loaded_at >= fresh).limit(5000).all())

    def get_prices(self, article):
        self.last_message = ""
        key = article_key(article)
        if not key:
            return []
        rows = [dict(item) for _, item in self._rows([key])]
        if not rows:
            self.last_message = ""  # пусто — не ошибка
        return rows

    def get_prices_many(self, articles):
        wanted = {}
        for article in articles or []:
            key = article_key(article)
            if key:
                wanted.setdefault(key, []).append(str(article))
        out = {}
        for key, item in self._rows(wanted):
            for alias in wanted.get(key, []):
                out.setdefault(alias, []).append(dict(item))
        return out

    def get_brand_candidates(self, article):
        seen, out = set(), []
        for item in self.get_prices(article):
            brand = str(item.get("brand") or "").strip()
            marker = (brand.upper(), str(item.get("provider") or ""))
            if brand and marker not in seen:
                seen.add(marker)
                out.append({"brand": brand, "article": item.get("article") or article, "name": item.get("name") or "",
                            "provider": item.get("provider") or self.name, "source": "прайс-лист"})
        return out

    def get_brands(self, article):
        return list(dict.fromkeys(c["brand"] for c in self.get_brand_candidates(article)))
